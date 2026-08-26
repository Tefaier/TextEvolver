import os
import random
import re
from io import BytesIO
from pathlib import Path

import bs4.element
import ebooklib
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.text.paragraph import CT_P
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from ebooklib import epub

from text_evolver.processing.binary_converter import convert_binary
from text_evolver.processing.browser import HTML_IMAGE_STYLE
from text_evolver.processing.images import get_image, get_pokemon_image

#from tqdm import tqdm
from text_evolver.processing.process_config_builder import ProcessingConfiguration, configure_process_unit
from text_evolver.processing.text_analysis import (
    convert_utf8_symbols,
    find_in_clean,
    possible_mutations,
    replace_iteration,
    string_with_meaning,
    text_cleaner,
)


def values_reset(object):
    object.settings = {}
    object.pokemons_list = {}
    object.units_list = {}
    object.word_conversions = {}
    object.direct_conversions = {}
    object.extra_img_list = {}
    object.word_counter = 0
    object.html_image_style = HTML_IMAGE_STYLE
    object.file_type = ''
    object.mutations_string = "|".join(possible_mutations)
    object.images_insert = []
    object.pokemons_navigation_map = []
    object.images_navigation_map = []
    object.feet_case = False
    object.clean_empty = False
    object.convert_to_utf = False


def delete_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    p._p = p._element = None


def image_choser(obj: dict): # if link is chosen return None else return binary
    official_image = obj.get("image_path")
    images_num = len(obj["binary"]) + (0 if official_image is None else 1)
    random_num = random.randint(0, images_num - 1)
    if official_image is not None and random_num==0:
        return None
    else:
        return obj["binary"][random_num - (0 if official_image is None else 1)]


def image_insert(object, unit, image_data: str):
    if object.file_type in ['fb2', 'epub']:
        try:
            number = str(object.images_insert.index(image_data)) + ".jpg"
        except:
            number = str(len(object.images_insert)) + ".jpg"
            object.images_insert.append(image_data)
            binary_tag = BeautifulSoup("", 'xml').new_tag('binary', id=number, **{'content-type': 'image/jpeg'})
            binary_tag.string = image_data
            unit.insert_before(binary_tag)
        new_tag = BeautifulSoup("", 'xml').new_tag('image', **{'href':"#" + number})
        unit.insert_before(new_tag)
    elif object.file_type == 'docx':
        pp = CT_P.add_p_before(unit._element)
        p = Paragraph(pp, unit._parent)
        #p = unit.insert_paragraph_before('')
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p.add_run().add_picture(BytesIO(convert_binary(image_data, "PIL")), width=Cm(12))
    elif object.file_type == 'html':
        new_tag = BeautifulSoup("", 'html.parser').new_tag('img', src="data:image/jpeg;base64," + image_data, style=HTML_IMAGE_STYLE)
        unit.insert_before(new_tag)


class ProcessUnit:
    settings: dict
    pokemons_list: dict
    units_list: dict
    word_conversions: dict
    direct_conversions: dict
    extra_img_list: dict

    images_insert: list
    pokemons_navigation_map: list
    images_navigation_map: list

    html_image_style: str
    file_type: str
    mutations_string: str
    word_counter: int

    #feet_case: bool
    #clean_empty: bool
    #convert_to_utf: bool

    def __init__(self, configuration: ProcessingConfiguration):
        values_reset(self)
        configure_process_unit(self, configuration)
        self.pokemons_navigation_map = [[x.lower(), x] for x in self.pokemons_list.keys()]
        self.images_navigation_map = [[x.lower(), x] for x in self.extra_img_list.keys()]

    def images_locate(self, string: str, unit):
        clean_text = text_cleaner(string, self.settings)
        if self.settings["pokemon"]:
            results = re.findall(fr"\b(?i:({'|'.join(self.pokemons_list.keys())}))({self.mutations_string})?\b", clean_text)
            for result in results:
                key = list(filter(lambda x: x[0] == result[0].lower(), self.pokemons_navigation_map))[0][1]
                item = self.pokemons_list[key]
                if (item["last word"] == None) or (item["last word"] + item["separation"] < self.word_counter and item["separation"] != 1):
                        image_data = get_pokemon_image(
                            key,
                            self.settings,
                            item["image_path"],
                            item["height"],
                            item["weight"],
                            image_choser(item),
                            item["explanation"],
                        )
                        if image_data != None:
                            self.pokemons_list[key]["last word"] = self.word_counter
                            image_insert(self, unit, image_data)
        if len(self.extra_img_list) != 0:
            results = re.findall(fr"\b(?i:({'|'.join(self.extra_img_list.keys())}))({self.mutations_string})?\b", clean_text)
            for result in results:
                key = list(filter(lambda x: x[0] == result[0].lower(), self.images_navigation_map))[0][1]
                item = self.extra_img_list[key]
                if (result[1] == '' or item["mutation"]) and ((item["last word"] == None) or (item["last word"] + item["separation"] < self.word_counter and item["separation"] != 1)):
                        image_data = get_image(image_choser(item), key, key, item["explanation"])
                        if image_data != None:
                            self.extra_img_list[key]["last word"] = self.word_counter
                            image_insert(self, unit, image_data)

    def direct_replace(self, string: str):  # direct replacement, convertation of symbols to utf
        for item in self.direct_conversions.items():
            string = string.replace(item[0], item[1])
        if self.settings["convert to utf"]:
            string = convert_utf8_symbols(string)
        return string

    def p_process(self, unit):  # processing of text unit for html
        if self.settings["clean empty"] and not string_with_meaning(unit.text):
            unit.decompose()
            return
        self.images_locate(unit.text, unit)
        parts = unit.contents
        for part in parts:
            text = self.direct_replace(part.text)
            words = text.split(' ')
            if words not in [[''], ['\n']]:
                words = self.text_alteration(words)
                self.word_counter += len(words)
                if type(part) == bs4.element.NavigableString:
                    part.replace_with(bs4.element.NavigableString(' '.join(words)))
                else:
                    part.string = ' '.join(words)

    def text_alteration(self, words: list):  # change last words for any
        clean_text = text_cleaner(" ".join(words), self.settings).split(" ")

        for item in self.units_list.items():
            results = find_in_clean(clean_text, item[1]["split"], item[0], False, {"units": item[1]["conversion"], "replace_with": item[1]["new unit"], "feet": (self.units_list['feet']["conversion"] if self.settings['feet check'] else None), "can be word": item[1]["can be word"]})
            if results["found"]:
                try:
                    words = replace_iteration(results["replace_map"], words.copy())
                    clean_text = text_cleaner(" ".join(words), self.settings).split(" ")
                except:
                    pass
        for item in self.word_conversions.items():
            results = find_in_clean(clean_text, item[1]["split"], item[0], item[1]["mutation"], {"units": None, "replace_with": item[1]["new words"], "feet": None, "can be word": None})
            if results["found"]:
                try:
                    words = replace_iteration(results["replace_map"], words.copy())
                    clean_text = text_cleaner(" ".join(words), self.settings).split(" ")
                except:
                    pass

        return words

    def p_process_word(self, par, table=None):  # processing of text unit for word
        if self.settings["clean empty"] and not string_with_meaning(par.text):
            delete_paragraph(par)
            return
        for run in par.runs:
            text = self.direct_replace(run.text)
            if (text not in ['', '\n', ' ']):
                self.images_locate(text, par if table is None else table)
                words = text.split(' ')
                words = self.text_alteration(words)
                self.word_counter += len(words)
                run.text = ' '.join(words)

    def remake_text(self, file_from, file_to):
        self.images_insert = []
        self.file_type = str(file_from).rsplit('.', 1)[-1].lower()
        if (self.file_type == "html"):  # html - check p AND span
            with open(file_from, "r", encoding="UTF-8") as file_read:
                soup = BeautifulSoup(file_read, 'html.parser')
            units = soup.find_all(['p', 'span'])
            #total_num = len(units)
            #tq = tqdm(total=total_num, bar_format='{l_bar}{bar:30}{r_bar}')
            for unit in units:
                if unit.find_parent('p' if unit.name == 'span' else 'span') is None:
                    self.p_process(unit)
                #tq.update(1)
            with open(file_to, "wb") as file:
                file.write(soup.encode())
            os.remove(file_from)
            #tq.close()
            return file_to

        elif (self.file_type == "epub"):  # epub
            file_read = epub.read_epub(file_from)
            items = file_read.get_items()
            #tq_main = tqdm(total=len(file_read.items), bar_format='{l_bar}{bar:20}{r_bar}', position=0)
            for item in items:
                if item.get_type() != ebooklib.ITEM_DOCUMENT:
                    continue
                soup = BeautifulSoup(item.get_content(), features='xml')
                units = soup.find_all(['p', 'span'])
                for unit in units:
                    if unit.find_parent('p' if unit.name == 'span' else 'span') is None:
                        self.p_process(unit)
                item.set_content(soup.encode())
                self.images_insert = []
                #tq_main.update(1)
            epub.write_epub(file_to, file_read)
            os.remove(file_from)
            #tq_main.close()
            return file_to

        elif (self.file_type == "docx"):  # docx
            document = Document(file_from)
            #total_num = len(document.paragraphs)
            #tq = tqdm(total=total_num, bar_format='{l_bar}{bar:30}{r_bar}')
            for paragraph in document.paragraphs:
                self.p_process_word(paragraph)
                #tq.update()
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            self.p_process_word(paragraph, table)
            document.save(file_to)
            os.remove(file_from)
            #tq.close()
            return file_to

        elif (self.file_type == "fb2"):  # fb2 - check p
            with open(file_from, "r", encoding="UTF-8") as file_read:
                soup = BeautifulSoup(file_read, 'xml')
            units = soup.find_all(['p', 'span'])
            #total_num = len(units)
            #tq = tqdm(total=total_num, bar_format='{l_bar}{bar:30}{r_bar}')
            for unit in units:
                if unit.find_parent('p' if unit.name == 'span' else 'span') is None:
                    self.p_process(unit)
                #tq.update(1)
            with open(file_to, "wb") as file:
                file.write(soup.encode())
            os.remove(file_from)
            #tq.close()
            return file_to


def process_files(
    configuration: ProcessingConfiguration,
    origin_directory: str | Path,
    output_directory: str | Path,
) -> int:
    '''Entry poing of processing files based on configuration'''
    origin = Path(origin_directory)
    output = Path(output_directory)
    if not origin.is_dir():
        raise FileNotFoundError(f"Input directory does not exist: {origin}")
    output.mkdir(parents=True, exist_ok=True)
    remaker = ProcessUnit(configuration)
    processed = 0
    for source in sorted(origin.iterdir()):
        if not source.is_file():
            continue
        remaker.remake_text(source, output / source.name)
        processed += 1
    return processed
