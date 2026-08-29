from io import BytesIO
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.text.paragraph import CT_P
from docx.shared import Cm
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from text_evolver.processing.binary_converter import convert_binary
from text_evolver.processing.documents.base import DocumentAdapter, DocumentPart


class DocxDocumentAdapter(DocumentAdapter):
    def __init__(self, source: Path, target: Path) -> None:
        super().__init__(source, target)
        self._document: DocumentObject = Document(source)
        self._blocks: list[tuple[Paragraph, Paragraph | Table]] = [
            (paragraph, paragraph) for paragraph in self._document.paragraphs
        ]
        self._blocks.extend(
            (paragraph, table)
            for table in self._document.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        )
        self._block_index = 0
        self._part_index = 0
        self._current_parts: list[Run | None] = []
        self._current_text = ""
        self._last_paragraph: Paragraph | None = None
        self._last_anchor: Paragraph | Table | None = None
        self._last_run: Run | None = None

    def read_part(self) -> DocumentPart | None:
        while True:
            if self._part_index >= len(self._current_parts):
                if self._block_index >= len(self._blocks):
                    self._clear_native_state()
                    return None
                paragraph, anchor = self._blocks[self._block_index]
                self._block_index += 1
                self._current_parts = list(paragraph.runs) or [None]
                self._part_index = 0
                self._current_text = paragraph.text
                self._last_paragraph = paragraph
                self._last_anchor = anchor
            part_index = self._part_index
            run = self._current_parts[part_index]
            self._part_index += 1
            self._last_run = run
            self._mark_part_read()
            return DocumentPart(
                text=run.text if run is not None else "",
                block_text=self._current_text,
                starts_block=part_index == 0,
            )

    def overwrite_last_part(self, text: str) -> None:
        self._require_last_part()
        if self._last_paragraph is None:
            raise RuntimeError("Document part has no containing paragraph")
        if self._last_run is None:
            if text:
                self._last_run = self._last_paragraph.add_run(text)
            return
        self._last_run.text = text

    def remove_last_block(self) -> None:
        self._require_last_part()
        if self._last_paragraph is None:
            raise RuntimeError("Document part has no containing paragraph")
        paragraph_element = self._last_paragraph._element
        paragraph_element.getparent().remove(paragraph_element)
        self._last_paragraph._p = self._last_paragraph._element = None
        self._part_index = len(self._current_parts)
        self._clear_native_state()

    def insert_image_before_last_block(self, image_data: str) -> None:
        self._require_last_part()
        if self._last_anchor is None:
            raise RuntimeError("Document part has no image insertion anchor")
        paragraph_element = CT_P.add_p_before(self._last_anchor._element)
        paragraph = Paragraph(paragraph_element, self._last_anchor._parent)
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        paragraph.add_run().add_picture(BytesIO(convert_binary(image_data, "PIL")), width=Cm(12))

    def _clear_native_state(self) -> None:
        self._last_paragraph = None
        self._last_anchor = None
        self._last_run = None
        self._clear_last_part()

    def _write(self, target: Path) -> None:
        self._document.save(target)
