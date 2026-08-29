import base64
import hashlib
from io import BytesIO
from pathlib import Path

import ebooklib
import pytest
from bs4 import BeautifulSoup
from docx import Document
from ebooklib import epub
from PIL import Image

from text_evolver.processing.documents import create_document_adapter


def jpeg_base64() -> str:
    output = BytesIO()
    Image.new("RGB", (8, 8), "red").save(output, format="JPEG")
    return base64.b64encode(output.getvalue()).decode("ascii")


def test_adapter_requires_a_last_read_part(tmp_path: Path):
    source = tmp_path / "source.html"
    target = tmp_path / "target.html"
    source.write_text("<p>text</p>", encoding="utf-8")
    adapter = create_document_adapter(source, target)

    with pytest.raises(RuntimeError, match="Read a document part"):
        adapter.overwrite_last_part("replacement")
    with pytest.raises(RuntimeError, match="Read a document part"):
        adapter.remove_last_block()
    with pytest.raises(RuntimeError, match="Read a document part"):
        adapter.insert_image_before_last_block(jpeg_base64())

    assert adapter.read_part() is not None
    assert adapter.read_part() is None
    with pytest.raises(RuntimeError, match="Read a document part"):
        adapter.overwrite_last_part("replacement")


def test_markup_remove_skips_remaining_parts_of_block(tmp_path: Path):
    source = tmp_path / "source.html"
    target = tmp_path / "target.html"
    source.write_text("<p><strong>remove</strong> tail</p><p>keep</p>", encoding="utf-8")
    adapter = create_document_adapter(source, target)

    first = adapter.read_part()
    assert first is not None and first.starts_block
    adapter.remove_last_block()
    second = adapter.read_part()
    assert second is not None and second.starts_block and second.text == "keep"
    adapter.overwrite_last_part("kept")
    adapter.save()

    soup = BeautifulSoup(target.read_bytes(), "html.parser")
    assert [value.get_text() for value in soup.find_all("p")] == ["kept"]
    assert not source.exists()


def test_empty_markup_and_docx_blocks_emit_sentinel_parts(tmp_path: Path):
    html_source = tmp_path / "empty.html"
    html_source.write_text("<p></p>", encoding="utf-8")
    html_adapter = create_document_adapter(html_source, tmp_path / "empty-output.html")

    html_part = html_adapter.read_part()
    assert html_part is not None
    assert (html_part.text, html_part.block_text, html_part.starts_block) == ("", "", True)

    docx_source = tmp_path / "empty.docx"
    document = Document()
    document.add_paragraph()
    document.save(docx_source)
    docx_adapter = create_document_adapter(docx_source, tmp_path / "empty-output.docx")

    docx_part = docx_adapter.read_part()
    assert docx_part is not None
    assert (docx_part.text, docx_part.block_text, docx_part.starts_block) == ("", "", True)


def test_failed_save_keeps_source_and_removes_partial_target(monkeypatch, tmp_path: Path):
    source = tmp_path / "source.html"
    target = tmp_path / "target.html"
    source.write_text("<p>text</p>", encoding="utf-8")
    adapter = create_document_adapter(source, target)

    def fail_write(path: Path) -> None:
        path.write_text("partial", encoding="utf-8")
        raise OSError("write failed")

    monkeypatch.setattr(adapter, "_write", fail_write)

    with pytest.raises(OSError, match="write failed"):
        adapter.save()

    assert source.exists()
    assert not target.exists()
    assert not (tmp_path / ".target.html.tmp").exists()


def test_unsupported_document_is_not_consumed(tmp_path: Path):
    source = tmp_path / "source.txt"
    source.write_text("text", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported document format"):
        create_document_adapter(source, tmp_path / "target.txt")

    assert source.exists()


def test_html_image_insertion(tmp_path: Path):
    source = tmp_path / "source.html"
    target = tmp_path / "target.html"
    source.write_text("<p>text</p>", encoding="utf-8")
    adapter = create_document_adapter(source, target)

    assert adapter.read_part() is not None
    adapter.insert_image_before_last_block(jpeg_base64())
    adapter.save()

    soup = BeautifulSoup(target.read_bytes(), "html.parser")
    assert len(soup.find_all("img")) == 1


def test_fb2_image_insertion(tmp_path: Path):
    source = tmp_path / "source.fb2"
    target = tmp_path / "target.fb2"
    source.write_text("<FictionBook><body><section><p>text</p></section></body></FictionBook>", encoding="utf-8")
    adapter = create_document_adapter(source, target)

    assert adapter.read_part() is not None
    image_data = jpeg_base64()
    adapter.insert_image_before_last_block(image_data)
    adapter.insert_image_before_last_block(image_data)
    assert adapter._inserted_images == [hashlib.sha256(image_data.encode("ascii")).digest()]
    adapter.save()

    soup = BeautifulSoup(target.read_bytes(), "xml")
    assert len(soup.find_all("binary")) == 1
    assert len(soup.find_all("image")) == 2


def test_docx_image_insertion(tmp_path: Path):
    source = tmp_path / "source.docx"
    target = tmp_path / "target.docx"
    document = Document()
    document.add_paragraph("text")
    document.save(source)
    adapter = create_document_adapter(source, target)

    assert adapter.read_part() is not None
    adapter.insert_image_before_last_block(jpeg_base64())
    adapter.save()

    processed = Document(target)
    assert len(processed.inline_shapes) == 1


def test_epub_image_insertion_preserves_existing_representation(tmp_path: Path):
    source = tmp_path / "source.epub"
    target = tmp_path / "target.epub"
    book = epub.EpubBook()
    book.set_identifier("fixture")
    book.set_title("Fixture")
    book.set_language("en")
    chapter = epub.EpubHtml(title="Chapter", file_name="chapter.xhtml", lang="en")
    chapter.content = "<html><body><p>text</p></body></html>"
    book.add_item(chapter)
    book.toc = (chapter,)
    book.spine = ["nav", chapter]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    epub.write_epub(source, book)
    adapter = create_document_adapter(source, target)

    assert adapter.read_part() is not None
    adapter.insert_image_before_last_block(jpeg_base64())
    adapter.save()

    processed = epub.read_epub(target)
    document_items = [
        item.get_content() for item in processed.get_items() if item.get_type() == ebooklib.ITEM_DOCUMENT
    ]
    assert any(b"<binary" in content and b"<image" in content for content in document_items)
