import random
import os
import time
import re
import shutil

import bs4.element

from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.text.paragraph import CT_P
from docx.text.paragraph import Paragraph

from ebooklib import epub
from bs4 import BeautifulSoup
from io import BytesIO
#from tqdm import tqdm

from settings_control import get_settings
from text_analysis import find_in_clean, text_cleaner, replace_iteration, possible_mutations, string_with_meaning, convert_utf8_symbols
from image_controller import get_image, get_pokemon_image
from binary_converter import convert_binary
from web_settings import html_image_style, pokemons_link, pokemons_link_list

from web_app.models import db, Thread
import web_app


def values_reset(object):
    object.settings = {}
    object.pokemons_list = {}
    object.pokemons_link = pokemons_link
    object.pokemons_link_list = pokemons_link_list
    object.units_list = {}
    object.word_conversions = {}
    object.direct_conversions = {}
    object.extra_img_list = {}
    object.word_counter = 0
    object.html_image_style = html_image_style
    object.file_type = ''
    object.mutations_string = "|".join(possible_mutations)
    object.images_insert = []
    object.pokemons_navigation_map = []
    object.images_navigation_map = []
    object.feet_case = False
    object.clean_empty = False
    object.convert_to_utf = False


def delete_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    p._p = p._element = None


def image_choser(obj: dict): # if link is chosen return None else return binary
    try:
        link = obj["link"]
    except:
        link = None
    images_num = len(obj["binary"]) + (0 if link==None else 1)
    random_num = random.randint(0, images_num - 1)
    if link!=None and random_num==0:
        return None
    else:
        return obj["binary"][random_num - (0 if link==None else 1)]


def image_insert(object, unit, image_data: str):
    if object.file_type in ['fb2', 'epub']:
        try:
            number = str(object.images_insert.index(image_data)) + ".jpg"
        except:
            number = str(len(object.images_insert)) + ".jpg"
            object.images_insert.append(image_data)
            binary_tag = BeautifulSoup("", 'xml').new_tag('binary', id=number, **{'content-type': 'image/jpeg'})
            binary_tag.string = image_data
            unit.insert_before(binary_tag)
        new_tag = BeautifulSoup("", 'xml').new_tag('image', **{'href':"#" + number})
        unit.insert_before(new_tag)
    elif object.file_type == 'docx':
        pp = CT_P.add_p_before(unit._element)
        p = Paragraph(pp, unit._parent)
        #p = unit.insert_paragraph_before('')
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p.add_run().add_picture(BytesIO(convert_binary(image_data, "PIL")), width=Cm(12))
    elif object.file_type == 'html':
        new_tag = BeautifulSoup("", 'html.parser').new_tag('img', src="data:image/jpeg;base64," + image_data, style=html_image_style)
        unit.insert_before(new_tag)


class ProcessUnit:
    settings: dict
    pokemons_list: dict
    units_list: dict
    word_conversions: dict
    direct_conversions: dict
    extra_img_list: dict

    images_insert: list
    pokemons_navigation_map: list
    images_navigation_map: list

    html_image_style: str
    file_type: str
    mutations_string: str
    pokemons_link: str
    pokemons_link_list: str

    word_counter: int

    #feet_case: bool
    #clean_empty: bool
    #convert_to_utf: bool

    def __init__(this, settings_id):
        values_reset(this)
        get_settings(this, settings_id)
        this.pokemons_navigation_map = [[x.lower(), x] for x in this.pokemons_list.keys()]
        this.images_navigation_map = [[x.lower(), x] for x in this.extra_img_list.keys()]

    def images_locate(this, string: str, unit):
        clean_text = text_cleaner(string, this.settings)
        if this.settings["pokemon"]:
            results = re.findall(fr"\b(?i:({'|'.join(this.pokemons_list.keys())}))({this.mutations_string})?\b", clean_text)
            for result in results:
                key = list(filter(lambda x: x[0] == result[0].lower(), this.pokemons_navigation_map))[0][1]
                item = this.pokemons_list[key]
                if (item["last word"] == None) or (item["last word"] + item["separation"] < this.word_counter and item["separation"] != 1):
                        image_data = get_pokemon_image(key, this.settings, item["link"], image_choser(item), item["explanation"])
                        if image_data != None:
                            this.pokemons_list[key]["last word"] = this.word_counter
                            image_insert(this, unit, image_data)
        if len(this.extra_img_list) != 0:
            results = re.findall(fr"\b(?i:({'|'.join(this.extra_img_list.keys())}))({this.mutations_string})?\b", clean_text)
            for result in results:
                key = list(filter(lambda x: x[0] == result[0].lower(), this.images_navigation_map))[0][1]
                item = this.extra_img_list[key]
                if (result[1] == '' or item["mutation"]) and ((item["last word"] == None) or (item["last word"] + item["separation"] < this.word_counter and item["separation"] != 1)):
                        image_data = get_image(image_choser(item), key, key, item["explanation"])
                        if image_data != None:
                            this.extra_img_list[key]["last word"] = this.word_counter
                            image_insert(this, unit, image_data)

    def direct_replace(this, string: str):  # direct replacement, convertation of symbols to utf
        for item in this.direct_conversions.items():
            string = string.replace(item[0], item[1])
        if this.settings["convert to utf"]:
            string = convert_utf8_symbols(string)
        return string

    def p_process(this, unit):  # processing of text unit for html
        if this.settings["clean empty"] and not string_with_meaning(unit.text):
            unit.decompose()
            return
        this.images_locate(unit.text, unit)
        parts = unit.contents
        for part_index, part in enumerate(parts):
            text = this.direct_replace(part.text)
            words = text.split(' ')
            if words not in [[''], ['\n']]:
                words = this.text_alteration(words)
                this.word_counter += len(words)
                if type(part) == bs4.element.NavigableString:
                    unit.contents[part_index] = bs4.element.NavigableString(' '.join(words))
                else:
                    part.string = ' '.join(words)

    def text_alteration(this, words: list):  # change last words for any
        clean_text = text_cleaner(" ".join(words), this.settings).split(" ")

        for item in this.units_list.items():
            results = find_in_clean(clean_text, item[1]["split"], item[0], False, {"units": item[1]["conversion"], "replace_with": item[1]["new unit"], "feet": (this.units_list['feet']["conversion"] if this.settings['feet check'] else None), "can be word": item[1]["can be word"]})
            if results["found"]:
                try:
                    words = replace_iteration(results["replace_map"], words.copy())
                    clean_text = text_cleaner(" ".join(words), this.settings).split(" ")
                except:
                    pass
        for item in this.word_conversions.items():
            results = find_in_clean(clean_text, item[1]["split"], item[0], item[1]["mutation"], {"units": None, "replace_with": item[1]["new words"], "feet": None, "can be word": None})
            if results["found"]:
                try:
                    words = replace_iteration(results["replace_map"], words.copy())
                    clean_text = text_cleaner(" ".join(words), this.settings).split(" ")
                except:
                    pass

        return words

    def p_process_word(this, par, table=None):  # processing of text unit for word
        if this.settings["clean empty"] and not string_with_meaning(par.text):
            delete_paragraph(par)
            return
        for run in par.runs:
            text = this.direct_replace(run.text)
            if (text not in ['', '\n', ' ']):
                this.images_locate(text, par if table is None else table)
                words = text.split(' ')
                words = this.text_alteration(words)
                this.word_counter += len(words)
                run.text = ' '.join(words)

    def remake_text(this, file_from, file_to):
        this.images_insert = []
        this.file_type = file_from.split('.')[-1]
        if (this.file_type == "html"):  # html - check p AND span
            with open(file_from, "r", encoding="UTF-8") as file_read:
                soup = BeautifulSoup(file_read, 'html.parser')
            units = soup.find_all(['p', 'span'])
            #total_num = len(units)
            #tq = tqdm(total=total_num, bar_format='{l_bar}{bar:30}{r_bar}')
            for unit in units:
                if unit.find_parent('p' if unit.name == 'span' else 'span') is None:
                    this.p_process(unit)
                #tq.update(1)
            with open(file_to, "wb") as file:
                file.write(soup.encode())
            os.remove(file_from)
            #tq.close()
            return file_to

        elif (this.file_type == "epub"):  # epub
            file_read = epub.read_epub(file_from)
            items = file_read.get_items()
            #tq_main = tqdm(total=len(file_read.items), bar_format='{l_bar}{bar:20}{r_bar}', position=0)
            for item in items:
                soup = BeautifulSoup(item.get_content(), features='xml')
                units = soup.find_all(['p', 'span'])
                for unit in units:
                    if unit.find_parent('p' if unit.name == 'span' else 'span') is None:
                        this.p_process(unit)
                item.set_content(soup.encode())
                this.images_insert = []
                #tq_main.update(1)
            epub.write_epub(file_to, file_read)
            os.remove(file_from)
            #tq_main.close()
            return file_to

        elif (this.file_type == "docx"):  # docx
            document = Document(file_from)
            #total_num = len(document.paragraphs)
            #tq = tqdm(total=total_num, bar_format='{l_bar}{bar:30}{r_bar}')
            for paragraph in document.paragraphs:
                this.p_process_word(paragraph)
                #tq.update()
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            this.p_process_word(paragraph, table)
            document.save(file_to)
            os.remove(file_from)
            #tq.close()
            return file_to

        elif (this.file_type == "fb2"):  # fb2 - check p
            with open(file_from, "r", encoding="UTF-8") as file_read:
                soup = BeautifulSoup(file_read, 'xml')
            units = soup.find_all(['p', 'span'])
            #total_num = len(units)
            #tq = tqdm(total=total_num, bar_format='{l_bar}{bar:30}{r_bar}')
            for unit in units:
                if unit.find_parent('p' if unit.name == 'span' else 'span') is None:
                    this.p_process(unit)
                #tq.update(1)
            with open(file_to, "wb") as file:
                file.write(soup.encode())
            os.remove(file_from)
            #tq.close()
            return file_to


def files_processing(settings_id, user_id, directory: str, new_dir):
    remaker = ProcessUnit(settings_id)
    directory = directory.replace('/', '\\')
    try:
        files = [directory + '\\' + x for x in os.listdir(directory)]
        try:
            if type(files) == str:
                file = str(files)
                file_to = '\\'.join(file.split('\\')[:-2] + [new_dir, file.split('\\')[-1]])
                remaker.remake_text(file, file_to)
            elif type(files) == list:
                for file in files:
                    file_to = '\\'.join(file.split('\\')[:-2] + [new_dir.replace('/', ''), file.split('\\')[-1]])
                    remaker.remake_text(file, file_to)
        except KeyboardInterrupt:
            with web_app.app.app_context():
                thread = db.session.query(Thread).filter_by(user_id=user_id).first()
                if thread is not None:
                    thread.waits = True
                    db.session.add(thread)
                    db.session.commit()
        except:
            print(f"Error: fail with processing {directory}")
            shutil.rmtree(directory)
            os.mkdir(directory)
    except:
        print(f"Error: not found directory {directory}")
    del remaker
    with web_app.app.app_context():
        thread = db.session.query(Thread).filter_by(user_id=user_id).first()
        if thread is not None:
            db.session.delete(thread)
            db.session.commit()




if __name__ == '__main__':
    from multiprocessing import Process
    settings_id = "1"
    path_source = "files_to_process\\"
    file_paths = [os.path.join(path_source, file) for file in os.listdir(path_source)]
    start_time = time.time()
    with web_app.app.app_context():
        t1 = Process(target=files_processing, args=[0, 0])
        new_thread = Thread(ident=t1.ident)
        db.session.add(new_thread)
        db.session.commit()
        t1.start()
    print(round(time.time() - start_time, 2))

